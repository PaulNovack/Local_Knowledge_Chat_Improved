#!/usr/bin/env python3
"""Import a TA/Petro location-master XLSX into normalized ta_ SQLite tables.
Requires Python >=3.10 and openpyxl. Does not modify the input workbook.
"""
import argparse
import hashlib
import json
import logging
import re
import sqlite3
import os
import tempfile
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen
from zipfile import ZipFile, BadZipFile
from datetime import date, datetime, time, timezone
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
DEFAULT_DB=ROOT/'database'/'chat_history.sqlite3'

SCHEMA='''
CREATE TABLE IF NOT EXISTS ta_imports(import_id INTEGER PRIMARY KEY,filename TEXT NOT NULL,file_sha256 TEXT NOT NULL,started_at TEXT NOT NULL,finished_at TEXT,status TEXT NOT NULL,row_count INTEGER,error TEXT);
CREATE TABLE IF NOT EXISTS ta_locations(site_id TEXT PRIMARY KEY,location_id TEXT UNIQUE,brand TEXT,state TEXT,city TEXT,name TEXT,directions TEXT,address TEXT,po_box TEXT,postal_code TEXT,latitude REAL,longitude REAL,gas_brand TEXT,first_import_id INTEGER NOT NULL,last_import_id INTEGER NOT NULL);
CREATE INDEX IF NOT EXISTS ta_locations_state ON ta_locations(state);
CREATE TABLE IF NOT EXISTS ta_facilities(facility_id INTEGER PRIMARY KEY,site_id TEXT NOT NULL REFERENCES ta_locations(site_id),kind TEXT NOT NULL,name TEXT NOT NULL,UNIQUE(site_id,kind,name));
CREATE TABLE IF NOT EXISTS ta_contacts(contact_id INTEGER PRIMARY KEY,facility_id INTEGER NOT NULL REFERENCES ta_facilities(facility_id),kind TEXT NOT NULL,number TEXT NOT NULL,extension TEXT NOT NULL DEFAULT '',raw_value TEXT NOT NULL,UNIQUE(facility_id,kind,number,extension));
CREATE TABLE IF NOT EXISTS ta_amenities(amenity_id INTEGER PRIMARY KEY,name TEXT NOT NULL UNIQUE);
CREATE TABLE IF NOT EXISTS ta_location_amenities(site_id TEXT NOT NULL REFERENCES ta_locations(site_id),amenity_id INTEGER NOT NULL REFERENCES ta_amenities(amenity_id),availability INTEGER,quantity REAL,raw_values_json TEXT NOT NULL,PRIMARY KEY(site_id,amenity_id));
CREATE TABLE IF NOT EXISTS ta_hours(hours_id INTEGER PRIMARY KEY,facility_id INTEGER NOT NULL REFERENCES ta_facilities(facility_id),rule_number INTEGER NOT NULL,opens_at TEXT,closes_at TEXT,closes_next_day INTEGER NOT NULL DEFAULT 0,is_24_hours INTEGER NOT NULL DEFAULT 0,is_closed INTEGER NOT NULL DEFAULT 0,parse_status TEXT NOT NULL,raw_text TEXT,UNIQUE(facility_id,rule_number));
CREATE TABLE IF NOT EXISTS ta_hour_days(hours_id INTEGER NOT NULL REFERENCES ta_hours(hours_id) ON DELETE CASCADE,weekday INTEGER NOT NULL CHECK(weekday BETWEEN 0 AND 6),PRIMARY KEY(hours_id,weekday));
CREATE TABLE IF NOT EXISTS ta_source_columns(sheet TEXT NOT NULL,column_number INTEGER NOT NULL,header TEXT NOT NULL,PRIMARY KEY(sheet,column_number));
CREATE TABLE IF NOT EXISTS ta_source_rows(site_id TEXT NOT NULL REFERENCES ta_locations(site_id),sheet TEXT NOT NULL,row_number INTEGER NOT NULL,import_id INTEGER NOT NULL,values_json TEXT NOT NULL,PRIMARY KEY(site_id,sheet));
CREATE TABLE IF NOT EXISTS ta_import_issues(issue_id INTEGER PRIMARY KEY,import_id INTEGER NOT NULL,site_id TEXT,sheet TEXT,row_number INTEGER,field TEXT,issue TEXT NOT NULL,raw_value TEXT);
CREATE VIEW IF NOT EXISTS ta_hours_grouped AS
SELECT f.site_id,f.kind,f.name,h.*,(SELECT group_concat(weekday,',') FROM
 (SELECT weekday FROM ta_hour_days d WHERE d.hours_id=h.hours_id ORDER BY weekday)) AS weekdays
FROM ta_facilities f JOIN ta_hours h USING(facility_id);
'''
CORE={'Site ID','Brand','Location ID','State','City','Location','Directions','Address','P.O. Box','Zipcode','Phone','Fax','Latitude','Longitude','Gas Brand','Full Service Restaurant','QSR(s)','In Bay Service Hours','Emergency Roadside Assistance Hours'}
QUANTITIES={'Service Bays','Service Pits','Truck Parking Spaces','Handicapped Parking','Showers','Handicapped Showers','Total Diesel Dispensers/Lanes','Def Bulk Dispensers/Lanes','Weigh Scale'}
DAY=r'(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday|Tues|Thurs|Mon|Tue|Wed|Thu|Fri|Sat|Sun)'
DAYMAP={'mon':0,'tue':1,'wed':2,'thu':3,'fri':4,'sat':5,'sun':6}


class DownloadForms(HTMLParser):
    def __init__(self):
        super().__init__();self.current=None;self.forms=[]
    def handle_starttag(self,tag,attrs):
        attrs=dict(attrs)
        if tag=='form':self.current={'action':attrs.get('action',''),'method':attrs.get('method','get').lower(),'text':''}
    def handle_data(self,data):
        if self.current is not None:self.current['text']+=data
    def handle_endtag(self,tag):
        if tag=='form' and self.current is not None:self.forms.append(self.current);self.current=None


def download_workbook(directory,page_url='https://www.ta-petro.com/location/',timeout=60):
    headers={'User-Agent':'TALocationWorkbookIngester/1.0'}
    with urlopen(Request(page_url,headers=headers),timeout=timeout) as response:
        page=response.read().decode('utf-8');base=response.url
    parser=DownloadForms();parser.feed(page)
    candidates=[f for f in parser.forms if 'download location data' in ' '.join(f['text'].lower().split())]
    if len(candidates)!=1:raise ValueError('Expected one Download Location Data form; TA website may have changed')
    form=candidates[0]
    if form['method']!='get':raise ValueError('Download form is no longer GET')
    endpoint=urljoin(base,form['action'])
    if urlparse(endpoint).scheme!='https' or urlparse(endpoint).hostname!=urlparse(base).hostname:
        raise ValueError('Unexpected download form destination')
    directory.mkdir(parents=True,exist_ok=True)
    temporary=None
    try:
        with urlopen(Request(endpoint,headers=headers),timeout=timeout) as response:
            disposition=response.headers.get('Content-Disposition','')
            match=re.search(r'filename="?([^";]+)',disposition,re.I)
            filename=Path(match[1].strip().replace('\\','/')).name if match else 'ta_location_master.xlsx'
            if not filename.lower().endswith('.xlsx'):filename='ta_location_master.xlsx'
            with tempfile.NamedTemporaryFile(dir=directory,suffix='.xlsx',delete=False) as target:
                temporary=Path(target.name);size=0
                while block:=response.read(1024*1024):
                    size+=len(block)
                    if size>100*1024*1024:raise ValueError('Workbook download exceeds 100 MiB')
                    target.write(block)
        with ZipFile(temporary) as archive:
            if 'xl/workbook.xml' not in archive.namelist():raise ValueError('Downloaded file is not an XLSX workbook')
            if archive.testzip():raise ValueError('Workbook archive is damaged')
        destination=directory/filename
        os.replace(temporary,destination);temporary=None
        logging.info('Downloaded %s from %s',destination,endpoint)
        return destination
    finally:
        if temporary is not None:temporary.unlink(missing_ok=True)


def now():return datetime.now(timezone.utc).isoformat()


def serial(value):
    if isinstance(value,(datetime,date,time)):return value.isoformat()
    return value


def identifier(cell):
    v=cell.value
    if v is None:return None
    if isinstance(v,(int,float)) and not isinstance(v,bool) and float(v).is_integer():
        fmt=cell.number_format
        return str(int(v)).zfill(len(fmt)) if re.fullmatch('0+',fmt or '') else str(int(v))
    return str(v).strip()


def clock(value):
    m=re.fullmatch(r'(\d{1,2}):(\d{2})\s*([AP])M',value.strip(),re.I)
    if not m or not 1<=int(m[1])<=12 or int(m[2])>59:raise ValueError('Invalid clock')
    return f'{int(m[1])%12+(12 if m[3].upper()=="P" else 0):02d}:{int(m[2]):02d}'


def days(label):
    result=set()
    for m in re.finditer('('+DAY+r')\s*-\s*('+DAY+')',label,re.I):
        a,b=DAYMAP[m[1][:3].lower()],DAYMAP[m[2][:3].lower()]
        result.update((a+i)%7 for i in range((b-a)%7+1))
    for m in re.finditer(DAY,label,re.I):result.add(DAYMAP[m[0][:3].lower()])
    return sorted(result)


def parse_hours(value):
    base={'opens_at':None,'closes_at':None,'closes_next_day':0,'is_24_hours':0,'is_closed':0,'parse_status':'unknown','days':[]}
    if value is None or not str(value).strip():return [dict(base,parse_status='missing')]
    value=str(value).strip()
    if re.fullmatch(r'24\s*hours|24/7(?:/365)?',value,re.I):return [dict(base,is_24_hours=1,parse_status='parsed',days=list(range(7)))]
    rules=[]
    for chunk in re.split(r'[,;]\s*(?='+DAY+')',value,flags=re.I):
        m=re.fullmatch(r'('+DAY+r'(?:\s*(?:-|&)\s*'+DAY+r')*)\s*[:,]?\s*(.*)',chunk.strip(),re.I)
        if not m:return [base]
        weekday=days(m[1]);tail=m[2].strip()
        if tail.lower()=='closed':rule=dict(base,is_closed=1,parse_status='parsed',days=weekday)
        elif re.fullmatch(r'24\s*hours|24/7(?:/365)?',tail,re.I):rule=dict(base,is_24_hours=1,parse_status='parsed',days=weekday)
        else:
            interval=re.fullmatch(r'(\d{1,2}:\d{2}\s*[AP]M)\s*(?:-|to)\s*(\d{1,2}:\d{2}\s*[AP]M)',tail,re.I)
            if not interval:return [base]
            try:start,end=clock(interval[1]),clock(interval[2])
            except ValueError:return [base]
            full_day=start==end=='00:00'
            # This workbook repeatedly uses midnight-to-midnight for full-day operation.
            # Preserve an explicit interpretation status rather than hiding that inference.
            if start==end and not full_day:return [base]
            rule=dict(base,opens_at=start,closes_at=end,closes_next_day=int(end<=start),is_24_hours=int(full_day),
                      parse_status='interpreted_24_hours' if full_day else 'parsed',days=weekday)
        rules.append(rule)
    grouped={}
    for rule in rules:
        key=tuple(rule[k] for k in base if k!='days')
        if key in grouped:grouped[key]['days']=sorted(set(grouped[key]['days']+rule['days']))
        else:grouped[key]=rule
    return list(grouped.values())


def facility(conn,site,kind,name):
    conn.execute('INSERT OR IGNORE INTO ta_facilities(site_id,kind,name) VALUES (?,?,?)',(site,kind,name))
    return conn.execute('SELECT facility_id FROM ta_facilities WHERE site_id=? AND kind=? AND name=?',(site,kind,name)).fetchone()[0]


def number(value):
    if value in (None,''):return None
    try:return float(value)
    except (ValueError,TypeError):raise ValueError(f'Expected numeric value, got {value!r}')


def add_contact(conn,fid,kind,value):
    if value is None:return
    for raw in re.split(r'\s*(?:;|\n|\s/\s)\s*',str(value)):
        if not raw.strip():continue
        ext=re.search(r'(?:ext\.?|x)\s*(\d+)\s*$',raw,re.I)
        digits=re.sub(r'\D','',raw[:ext.start()] if ext else raw)
        normalized='+1'+digits if len(digits)==10 else '+'+digits if len(digits)==11 and digits[0]=='1' else raw.strip()
        conn.execute('INSERT OR IGNORE INTO ta_contacts(facility_id,kind,number,extension,raw_value) VALUES (?,?,?,?,?)',(fid,kind,normalized,ext[1] if ext else '',raw))


def ingest(path,db,sheet=None):
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    db.parent.mkdir(parents=True,exist_ok=True)
    with sqlite3.connect(db) as conn:
        conn.execute('PRAGMA foreign_keys=ON');conn.execute('PRAGMA journal_mode=WAL');conn.executescript(SCHEMA)
        run=conn.execute("INSERT INTO ta_imports(filename,file_sha256,started_at,status) VALUES (?,?,?,'running')",(path.name,digest,now())).lastrowid
        conn.commit();workbook=None
        try:
            workbook=load_workbook(path,read_only=True,data_only=False)
            sheets=[workbook[sheet]] if sheet else workbook.worksheets
            prepared=[];seen_sites=set();seen_locations=set()
            for ws in sheets:
                iterator=ws.iter_rows();header=next(iterator,None)
                if not header:continue
                headers=[str(c.value).strip() if c.value is not None else f'Column {i}' for i,c in enumerate(header,1)]
                required={'Site ID','Location ID','Location','Brand','Address','State','City','Latitude','Longitude'}
                if not required.issubset(headers):
                    if sheet:raise ValueError(f'{ws.title}: missing required columns {required-set(headers)}')
                    logging.info('Skipping non-location sheet %s',ws.title);continue
                for rownum,cells in enumerate(iterator,2):
                    if all(c.value is None for c in cells):continue
                    if any(c.data_type=='f' for c in cells):raise ValueError(f'{ws.title} row {rownum}: formula cells require materialized values before import')
                    values=[serial(c.value) for c in cells]
                    pairs=list(zip(headers,values));d={h:v for h,v in pairs}
                    site=identifier(cells[headers.index('Site ID')]);lid=identifier(cells[headers.index('Location ID')])
                    if not site:raise ValueError(f'{ws.title} row {rownum}: missing Site ID')
                    if site in seen_sites or (lid is not None and lid in seen_locations):raise ValueError(f'Duplicate Site ID or Location ID at {ws.title}:{rownum}')
                    seen_sites.add(site);seen_locations.add(lid)
                    lat,lon=number(d['Latitude']),number(d['Longitude'])
                    if lat is not None and not -90<=lat<=90:raise ValueError(f'{site}: invalid latitude')
                    if lon is not None and not -180<=lon<=180:raise ValueError(f'{site}: invalid longitude')
                    d['Zipcode']=identifier(cells[headers.index('Zipcode')]) if 'Zipcode' in headers else None
                    prepared.append((ws.title,rownum,headers,values,pairs,d,site,lid,lat,lon))
            if not prepared:raise ValueError('No location rows found; database not updated')
            with conn:
                for sheetname,rownum,headers,values,pairs,d,site,lid,lat,lon in prepared:
                    for i,h in enumerate(headers,1):conn.execute('INSERT INTO ta_source_columns VALUES (?,?,?) ON CONFLICT(sheet,column_number) DO UPDATE SET header=excluded.header',(sheetname,i,h))
                    conn.execute('''INSERT INTO ta_locations VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(site_id) DO UPDATE SET location_id=excluded.location_id,brand=excluded.brand,state=excluded.state,
                    city=excluded.city,name=excluded.name,directions=excluded.directions,address=excluded.address,po_box=excluded.po_box,
                    postal_code=excluded.postal_code,latitude=excluded.latitude,longitude=excluded.longitude,gas_brand=excluded.gas_brand,last_import_id=excluded.last_import_id''',
                    (site,lid,d.get('Brand'),d.get('State'),d.get('City'),d.get('Location'),d.get('Directions'),d.get('Address'),d.get('P.O. Box'),d.get('Zipcode'),lat,lon,d.get('Gas Brand'),run,run))
                    for t in ['ta_contacts','ta_hours']:
                        conn.execute(f'DELETE FROM {t} WHERE facility_id IN (SELECT facility_id FROM ta_facilities WHERE site_id=?)',(site,))
                    conn.execute('DELETE FROM ta_location_amenities WHERE site_id=?',(site,))
                    store=facility(conn,site,'store',str(d.get('Location') or site));active={store}
                    if lid is None:conn.execute('INSERT INTO ta_import_issues(import_id,site_id,sheet,row_number,field,issue) VALUES (?,?,?,?,?,?)',(run,site,sheetname,rownum,'Location ID','Missing source Location ID; retained using Site ID'))
                    add_contact(conn,store,'phone',d.get('Phone'));add_contact(conn,store,'fax',d.get('Fax'))
                    for field,kind in [('Full Service Restaurant','restaurant_full_service'),('QSR(s)','restaurant_quick_service')]:
                        for name in re.split(r'\s*[,;]\s*',str(d.get(field) or '')):
                            if name:active.add(facility(conn,site,kind,name))
                    for field in ['In Bay Service Hours','Emergency Roadside Assistance Hours']:
                        fid=facility(conn,site,'service',field.removesuffix(' Hours'));active.add(fid)
                        raw=d.get(field)
                        rules=parse_hours(raw)
                        for i,r in enumerate(rules,1):
                            hid=conn.execute('INSERT INTO ta_hours(facility_id,rule_number,opens_at,closes_at,closes_next_day,is_24_hours,is_closed,parse_status,raw_text) VALUES (?,?,?,?,?,?,?,?,?)',
                            (fid,i,r['opens_at'],r['closes_at'],r['closes_next_day'],r['is_24_hours'],r['is_closed'],r['parse_status'],None if raw is None else str(raw))).lastrowid
                            conn.executemany('INSERT INTO ta_hour_days VALUES (?,?)',[(hid,day) for day in r['days']])
                        if any(r['parse_status']=='unknown' for r in rules):conn.execute('INSERT INTO ta_import_issues(import_id,site_id,sheet,row_number,field,issue,raw_value) VALUES (?,?,?,?,?,?,?)',(run,site,sheetname,rownum,field,'Unparsed schedule',str(raw)))
                    for fid, in conn.execute('SELECT facility_id FROM ta_facilities WHERE site_id=?',(site,)).fetchall():
                        if fid not in active:conn.execute('DELETE FROM ta_facilities WHERE facility_id=?',(fid,))
                    features={}
                    for h,v in pairs:
                        if h not in CORE:features.setdefault(h,[]).append(v)
                    for h,vals in features.items():
                        nonempty=[v for v in vals if v not in (None,'')]
                        distinct={json.dumps(v,ensure_ascii=False) for v in nonempty}
                        conflicting=len(distinct)>1
                        if conflicting:conn.execute('INSERT INTO ta_import_issues(import_id,site_id,sheet,row_number,field,issue,raw_value) VALUES (?,?,?,?,?,?,?)',(run,site,sheetname,rownum,h,'Duplicate source columns disagree',json.dumps(vals)))
                        v=nonempty[0] if nonempty else None
                        flag=None;quantity=None
                        if not conflicting:
                            if h in QUANTITIES:
                                quantity=number(v);flag=int(quantity>0) if quantity is not None else None
                            elif v is not None and h=='Truck Wash':
                                flag=1
                            elif v is not None:
                                normalized=str(v).strip().lower()
                                flag=1 if normalized in ('y','yes','true','1') else 0 if normalized in ('n','no','false','0') else None
                                if flag is None:conn.execute('INSERT INTO ta_import_issues(import_id,site_id,sheet,row_number,field,issue,raw_value) VALUES (?,?,?,?,?,?,?)',(run,site,sheetname,rownum,h,'Unrecognized availability flag',str(v)))
                        conn.execute('INSERT OR IGNORE INTO ta_amenities(name) VALUES (?)',(h,))
                        aid=conn.execute('SELECT amenity_id FROM ta_amenities WHERE name=?',(h,)).fetchone()[0]
                        conn.execute('INSERT INTO ta_location_amenities VALUES (?,?,?,?,?)',(site,aid,flag,quantity,json.dumps(vals,ensure_ascii=False)))
                    conn.execute('INSERT INTO ta_source_rows VALUES (?,?,?,?,?) ON CONFLICT(site_id,sheet) DO UPDATE SET row_number=excluded.row_number,import_id=excluded.import_id,values_json=excluded.values_json',(site,sheetname,rownum,run,json.dumps(values,ensure_ascii=False)))
                conn.execute("UPDATE ta_imports SET status='complete',finished_at=?,row_count=? WHERE import_id=?",(now(),len(prepared),run))
            logging.info('Imported %d locations; run=%d; review issues=%d',len(prepared),run,conn.execute('SELECT count(*) FROM ta_import_issues WHERE import_id=?',(run,)).fetchone()[0])
            return len(prepared)
        except Exception as exc:
            with conn:conn.execute("UPDATE ta_imports SET status='failed',finished_at=?,error=? WHERE import_id=?",(now(),str(exc),run))
            raise
        finally:
            if workbook:workbook.close()


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('workbook',type=Path,nargs='?',help='Optional local workbook; omitted = download current file from TA/Petro')
    p.add_argument('--db',type=Path,default=DEFAULT_DB);p.add_argument('--sheet')
    p.add_argument('--download-dir',type=Path,default=Path('data/downloads'))
    p.add_argument('--page-url',default='https://www.ta-petro.com/location/')
    a=p.parse_args();logging.basicConfig(level=logging.INFO,format='%(levelname)s %(message)s')
    try:ingest(a.workbook or download_workbook(a.download_dir,a.page_url),a.db,a.sheet)
    except Exception:logging.exception('Import failed; existing location records preserved');return 1
    return 0

if __name__=='__main__':raise SystemExit(main())
